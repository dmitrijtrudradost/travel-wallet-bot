from pathlib import Path

from wallet.app import REPLY_MENU, App
from wallet.service import WalletService
from wallet.storage import Storage


class FakeRates:
    def __init__(self):
        self.quotes = {"USDRUB": 80, "USDCNY": 7.2}
        self.calls = []

    def fetch_quotes(self, currencies):
        self.calls.append(list(currencies))
        return self.quotes


def test_trip_flow_expense_and_history(tmp_path):
    rates = FakeRates()
    app = App(WalletService(Storage(tmp_path / "wallet.db"), rates, tmp_path))

    start = app.on_start(7, "Y Lover")
    assert "Привет, Y Lover!" in start.text
    assert "Что я умею" in start.text
    assert start.reply_keyboard is True
    assert start.buttons == []

    app.on_text(7, "🆕 Создать путешествие")
    app.on_text(7, "Россия")
    named = app.on_text(7, "Китай")
    assert "назвать поездку" in named.text.lower()
    assert rates.calls == []
    offer = app.on_text(7, "Китай 2026")
    assert "1 RUB = 0.0900 CNY" in offer.text
    assert rates.calls == [["RUB", "CNY"]]

    asked = app.on_callback(7, "rate:yes")
    assert "RUB" in asked.text
    created = app.on_text(7, "8000")
    assert "Путешествие создано" in created.text
    assert "Китай 2026" in created.text
    assert "720.00 CNY" in created.text

    preview = app.on_text(7, "100 юаней")
    assert "100.00 CNY = 1 111.11 RUB" in preview.text
    assert "Учесть как расход?" in preview.text
    asked_type = app.on_callback(7, "expense:yes")
    assert "тип расхода" in asked_type.text.lower()
    asked_receipt = app.on_text(7, "еда")
    assert "чек" in asked_receipt.text.lower()
    done = app.on_callback(7, "receipt:skip")
    assert "620.00 CNY" in done.text
    assert "6 888.89 RUB" in done.text
    assert done.buttons == []

    history = app.on_command(7, "history")
    assert history.file_path
    assert "еда" in history.text
    assert "Китай 2026" in history.text
    assert history.buttons
    picked = app.on_callback(7, history.buttons[-1][0][1])
    assert "Китай 2026" in picked.text
    assert len(rates.calls) == 1


def test_rub_to_rub_and_manual_rate(tmp_path):
    app = App(WalletService(Storage(tmp_path / "wallet.db"), FakeRates(), tmp_path))
    app.on_command(3, "newtrip")
    app.on_text(3, "Россия")
    asked = app.on_text(3, "Россия")
    assert "назвать поездку" in asked.text.lower()
    offer = app.on_text(3, "Дом")
    assert "одном валютном пространстве RUB" in offer.text
    assert "1 RUB = 1.0000 RUB" in offer.text

    manual = app.on_callback(3, "rate:no")
    assert "курс" in manual.text.lower()
    app.on_text(3, "1")
    created = app.on_text(3, "1000")
    assert "1 000.00 RUB" in created.text

    preview = app.on_text(3, "100")
    assert "100.00 RUB = 100.00 RUB" in preview.text


def test_unknown_country_asks_for_code(tmp_path):
    app = App(WalletService(Storage(tmp_path / "wallet.db"), FakeRates(), tmp_path))
    app.on_text(4, "🆕 Создать путешествие")
    asked = app.on_text(4, "Атлантида")
    assert "код" in asked.text.lower()
    next_step = app.on_text(4, "EUR")
    assert "страну" in next_step.text.lower()


def test_receipt_file_is_stored_and_linked(tmp_path):
    from openpyxl import load_workbook

    service = WalletService(Storage(tmp_path / "wallet.db"), FakeRates(), tmp_path)
    app = App(service)
    service.create_trip(7, "Россия", "Китай", "RUB", "CNY", 80 / 7.2, 8000, title="Китай 2026")
    app.on_text(7, "100")
    app.on_callback(7, "expense:yes")
    app.on_text(7, "еда")
    rejected = app.on_receipt(7, b"not-a-receipt", ".exe")
    assert "PDF" in rejected.text
    done = app.on_receipt(7, b"%PDF-1.4", ".pdf")
    assert "Чек сохранён" in done.text
    stored = service.storage.expenses(7)[0]["receipt_path"]
    assert stored.endswith(".pdf")
    assert Path(stored).read_bytes().startswith(b"%PDF")
    history = app.on_command(7, "history")
    cell = load_workbook(history.file_path).active.cell(2, 10)
    assert cell.value == "Чек"
    assert cell.hyperlink.target.endswith(".pdf")


def test_menu_labels():
    labels = [label for row in REPLY_MENU for label in row]
    assert labels == [
        "🆕 Создать путешествие",
        "🧳 Мои путешествия",
        "💰 Баланс",
        "📊 История расходов",
        "💱 Изменить курс",
    ]
