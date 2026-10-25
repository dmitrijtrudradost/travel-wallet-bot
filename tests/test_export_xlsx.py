from openpyxl import load_workbook

from wallet.export_xlsx import HEADERS, write_history


def test_history_file_has_date_time_and_type(tmp_path):
    path = write_history(
        tmp_path / "history_1.xlsx",
        [["25.10.2026", "16:40", "расход", "Китай 2026", "Россия → Китай", 100, "CNY", 1280, "RUB"]],
    )
    sheet = load_workbook(path).active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet[2][0].value == "25.10.2026"
    assert sheet[2][1].value == "16:40"
    assert sheet[2][2].value == "расход"


def test_history_file_links_receipt(tmp_path):
    receipt = tmp_path / "check.pdf"
    receipt.write_bytes(b"%PDF")
    path = write_history(
        tmp_path / "history_1.xlsx",
        [["25.10.2026", "16:40", "еда", "Китай", "Россия → Китай", 100, "CNY", 1280, "RUB", str(receipt)]],
    )
    cell = load_workbook(path).active.cell(2, 10)
    assert cell.value == "Чек"
    assert cell.hyperlink.target.endswith("check.pdf")
