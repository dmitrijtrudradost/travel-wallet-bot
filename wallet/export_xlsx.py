from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = [
    "Дата",
    "Время",
    "Тип траты",
    "Название",
    "Путешествие",
    "Сумма",
    "Валюта",
    "В домашней валюте",
    "Домашняя валюта",
    "Чек",
]


def write_history(path: Path, rows: list[list]) -> Path:
    """Сохраняет историю расходов в xlsx. Первая строка — заголовки."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "История расходов"
    sheet.append(HEADERS)
    link_font = Font(color="0563C1", underline="single")
    receipt_column = len(HEADERS)
    for row in rows:
        values = list(row)
        receipt = values[receipt_column - 1] if len(values) >= receipt_column else None
        values = values[: receipt_column - 1]
        values.extend([None] * (receipt_column - 1 - len(values)))
        values.append(None)
        sheet.append(values)
        if not receipt:
            continue
        cell = sheet.cell(sheet.max_row, receipt_column)
        cell.value = "Чек"
        cell.hyperlink = Path(str(receipt)).resolve().as_uri()
        cell.font = link_font
    workbook.save(path)
    return path
